#!/usr/bin/env python3
"""Export the frozen T+1 research results as a workbook and static figures.

Optional local research dependencies: matplotlib, openpyxl. No live application
imports this script.
"""
import csv
import gzip
import json
import shutil
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "research/limit-up-t1-2026"
ARTIFACTS = ROOT / "docs/assets/t1-2026"


def load(path):
    if path.suffix == ".gz":
        with gzip.open(path, "rt") as stream:
            return json.load(stream)
    return json.loads(path.read_text())


def main():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib import font_manager
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    ARTIFACTS.mkdir(parents=True, exist_ok=True)
    summary = load(DATA / "data-summary.json")
    rows = load(DATA / "replay.json.gz")
    examples = load(DATA / "transaction-evidence-examples.json")
    selected_ids = summary["selection"]["selection"]["rule"]
    from banxia_strategy.t1_research import matches
    selected = [r for r in rows if matches(r, selected_ids)]
    for path in ("/System/Library/Fonts/PingFang.ttc", "/System/Library/Fonts/STHeiti Light.ttc"):
        if Path(path).exists():
            font_manager.fontManager.addfont(path)
            plt.rcParams["font.family"] = font_manager.FontProperties(fname=path).get_name()
            break
    plt.rcParams.update({"axes.spines.top": False, "axes.spines.right": False,
                         "axes.unicode_minus": False, "font.size": 11, "figure.dpi": 160})

    monthly = summary["monthly"]
    fig, ax = plt.subplots(figsize=(10, 4.8), layout="constrained")
    positions = list(range(len(monthly)))
    ax.bar(positions, [r["success_rate"] * 100 for r in monthly], color="#236b8e", width=.6)
    ax.errorbar(positions, [r["success_rate"] * 100 for r in monthly],
                yerr=[[100 * (r["success_rate"] - r["wilson_lower"]) for r in monthly],
                      [100 * (r["wilson_upper"] - r["success_rate"]) for r in monthly]],
                fmt="none", ecolor="#20303d", capsize=4, lw=1.2)
    for i, r in enumerate(monthly):
        ax.text(i, 3, f"n={r['buys']}", ha="center", fontsize=9, color="white")
    ax.axhline(80, color="#c34738", linestyle="--", label="目标：80%")
    ax.set(xticks=positions, xticklabels=[r["month"][5:] + "月" for r in monthly],
           ylabel="扣费收益 >5% 的比例（%）", ylim=(0, 100),
           title="主板首板基准：月度达标率与 Wilson 95% 区间")
    ax.legend(frameon=False, loc="upper right")
    fig.savefig(ARTIFACTS / "monthly.png")
    plt.close(fig)

    stages = [r for r in summary["selected_performance"] if r["stage"] in ("train", "validation", "holdout")]
    labels = {"train": "1—6月训练", "validation": "7—8月验证", "holdout": "9月留出"}
    fig, ax = plt.subplots(figsize=(9, 4.8), layout="constrained")
    for i, row in enumerate(stages):
        value = row["success_rate"] * 100
        ax.errorbar(value, i, xerr=[[value - row["wilson_lower"] * 100],
                                   [row["wilson_upper"] * 100 - value]], fmt="o",
                    color="#236b8e" if i < 2 else "#b36b24", capsize=6, markersize=9)
        ax.text(value, i + .22, f"{row['successes']}/{row['buys']} = {value:.1f}%",
                va="center", ha="center")
    ax.axvline(80, color="#c34738", linestyle="--")
    ax.set(yticks=range(3), yticklabels=[labels[r["stage"]] for r in stages],
           xlim=(0, 100), ylim=(-.5, 2.6), xlabel="扣费收益 >5% 的比例（%）",
           title=f"冻结规则 {selected_ids['id']}：达标率与 Wilson 95% 区间")
    fig.savefig(ARTIFACTS / "validation.png")
    plt.close(fig)

    base = [r for r in rows if r["entry_mode"] == "early" and r["board"] == "main"
            and r["consecutive_limit_days"] == 1 and (r["amount_cny"] or 0) >= 200000000]
    values = [r["target_net_return_pct"] for r in base if "target_net_return_pct" in r]
    fig, ax = plt.subplots(figsize=(9, 4.8), layout="constrained")
    ax.hist(values, bins=40, color="#236b8e", edgecolor="white", linewidth=.4)
    ax.axvline(5, color="#c34738", linestyle="--", label="收益严格超过5%")
    ax.set(xlabel="单笔扣费收益（%）", ylabel="交易次数",
           title="主板首板基准：预设止盈及尾盘退出的收益分布")
    ax.legend(frameon=False)
    fig.savefig(ARTIFACTS / "returns.png")
    plt.close(fig)

    columns = {
        "trade_date": "D1涨停日期", "symbol": "代码", "name": "名称", "board": "板块",
        "industry": "当前行业", "consecutive_limit_days": "D1连板数",
        "amount_cny": "D1成交额元", "total_market_cap_cny": "D1总市值元",
        "float_market_cap_cny": "D1流通市值元", "turnover_pct": "D1换手率%",
        "amplitude_pct": "D1振幅%", "open_change_pct": "D1开盘涨幅%",
        "change_pct": "D1涨幅%", "return_5d_pct": "D1近5日涨幅%",
        "capital_basis": "股本依据", "buy_date": "D2买入日期", "d2_open_gap_pct": "D2开盘涨幅%",
        "entry_mode": "入场模式", "buy_status": "买入状态", "signal_time": "信号时间",
        "buy_time": "买入分钟结束", "buy_sample_price": "买入分时价格", "buy_price": "买入模型价格",
        "shares": "股数", "buy_minute_volume_hands": "买入分钟成交手数",
        "buy_participation_pct": "分钟成交占比%", "sell_date": "D3卖出日期",
        "open_benchmark_net_return_pct": "D3开盘基准净收益%",
        "early_sell_status": "早盘卖出状态", "early_sell_time": "早盘卖出分钟结束",
        "early_sell_price": "早盘卖出模型价格", "early_net_return_pct": "早盘卖出净收益%",
        "target_sell_status": "止盈卖出状态", "target_sell_time": "止盈或尾盘卖出分钟结束",
        "target_sell_price": "止盈或尾盘模型价格", "target_net_return_pct": "止盈或尾盘净收益%",
        "target_sell_reason": "卖出原因", "split": "样本阶段",
    }
    winners = sorted([r for r in rows if any(r.get(f"{mode}_net_return_pct", -100) > 5
                                           for mode in ("early", "target"))],
                     key=lambda r: (r["trade_date"], r["symbol"], r["entry_mode"]))
    with (DATA / "达标股票明细.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream, lineterminator="\n")
        writer.writerow(columns.values())
        writer.writerows([[r.get(k) for k in columns] for r in winners])
    wb = Workbook()
    notes = wb.active
    notes.title = "口径与说明"
    notes.append(["项目", "说明"])
    for row in (
        ("目标", "D1涨停；D2 10点前买入；D3卖出；扣费收益严格超过5%"),
        ("统计截止", "D3不晚于2026-09-30；最后可完整复盘D1为2026-09-28"),
        ("成交证据", "分钟价格/成交量代理；12例另有历史分笔佐证；不能证明个人排队成交"),
        ("重复口径", "同一股票日期可有early/breakout/reclaim三种固定买点，不能把行数当去重事件数"),
        ("当前名称行业", "名称和行业为当前mootdx分类；不代表历史题材或历史ST状态"),
        ("费用", "每笔预算1万元，双边佣金万3最低5元、过户费万0.1、卖出印花税万5、双边滑点0.1%"),
        ("止盈", "按净收益至少5.1%向上取整至分，满足价格/成交量证据后按预定价退出"),
        ("没有卖出", "已买入但D3数据未知/未成交仍放入保守成功率分母"),
        ("规则状态", "研究规则未激活；任何历史比例都不保证未来80%"),
    ):
        notes.append(row)

    def add_sheet(title, records, fields):
        sheet = wb.create_sheet(title)
        sheet.append(list(fields.values()))
        for row in records:
            sheet.append([row.get(key) for key in fields])
        sheet.freeze_panes = "C2"
        sheet.auto_filter.ref = sheet.dimensions
        return sheet

    add_sheet("达标股票明细", winners, columns)
    add_sheet("冻结规则全部样本", sorted(selected, key=lambda r: (r["buy_date"], r["symbol"])), columns)
    add_sheet("分笔核验样例", examples, {
        **columns, "buy_transaction_bucket": "买入分笔时间段",
        "buy_eligible_transaction_hands": "买价内实际成交手数",
        "buy_transaction_capacity_corroborated": "分笔买入容量佐证",
        "sell_eligible_transaction_hands": "卖价内实际成交手数",
        "sell_transaction_capacity_corroborated": "分笔卖出容量佐证",
    })
    add_sheet("月度基准", monthly, {
        "month": "月份", "buys": "买入代理次数", "successes": "收益大于5%次数",
        "success_rate": "达标率", "unknown_or_unfilled_exits": "卖出未知或未成交",
        "mean_known_return_pct": "已知卖出平均收益%", "worst_known_return_pct": "最差已知收益%",
    })
    for sheet in wb:
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="236B8E")
        for i in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(i)].width = 19
        sheet.row_dimensions[1].height = 30
    notes.column_dimensions["B"].width = 110
    wb.save(DATA / "2026涨停隔日交易复盘.xlsx")
    chart_spec = {"tool": "generate_line_chart", "args": {
        "data": [{"time": r["month"], "value": round(r["success_rate"] * 100, 4)}
                 for r in monthly],
        "title": "主板首板基准：月度净收益超过5%的比例",
        "axisXTitle": "D2买入月份", "axisYTitle": "达标率（%）",
        "width": 1000, "height": 500, "theme": "academy",
    }}
    (DATA / "monthly-chart-spec.json").write_text(json.dumps(chart_spec, ensure_ascii=False, indent=2))
    for name in (
        "2026涨停隔日交易复盘.xlsx", "达标股票明细.csv", "selected-trades.csv",
        "data-summary.json", "collection-manifest.json", "frozen-selection.json",
        "transaction-evidence-examples.csv", "source-clock-audit.json.gz",
        "monthly-chart-spec.json", "feature-comparison.csv", "monthly.csv",
    ):
        if name.endswith(".csv"):
            (ARTIFACTS / name).write_text((DATA / name).read_text(), encoding="utf-8")
        else:
            shutil.copy2(DATA / name, ARTIFACTS / name)
    original = DATA / "audit-before-star-lot-correction"
    if original.exists():
        old = load(original / "data-summary.json")
        assert old["selected_performance"] == summary["selected_performance"]
        assert old["selection"]["selection"]["rule"] == summary["selection"]["selection"]["rule"]
        shutil.copy2(original / "frozen-selection.json", ARTIFACTS / "original-v2-selection.json")
        shutil.copy2(original / "protocol-v2.json", ARTIFACTS / "original-v2-protocol.json")
    print(json.dumps({"winner_rows": len(winners), "workbook": str(DATA / "2026涨停隔日交易复盘.xlsx"),
                      "figures": str(ARTIFACTS)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
